"""Downloads folder lookup, Daily Summary exports, and backup extras not covered elsewhere."""
import os
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook

from tests.helpers import ServiceTestCase
from vx7khata import exports, paths


class TestDownloadsDir(ServiceTestCase):
    def test_posix_downloads_under_home(self):
        with mock.patch.object(paths.os, "name", "posix"):
            self.assertEqual(paths.downloads_dir(), Path.home() / "Downloads")


class TestDailyExports(ServiceTestCase):
    def setUp(self):
        super().setUp()
        a = self.svc.add_customer("Amit")
        b = self.svc.add_customer("Bharat")
        self.add(a.id, "UDHAAR", "03-10-2026", "Chain", "5000")
        self.add(a.id, "JAMA", "03-10-2026", "Cash", "1000")
        self.add(b.id, "UDHAAR", "03-10-2026", "Ring", "2500")
        self.s = self.svc.daily_summary("2026-10-03")

    def test_xlsx_matches_summary(self):
        p = exports.export_daily_summary_xlsx(self.s, self.s.rows, Path(self.tmp.name) / "d.xlsx", "SHREEJI BOXES")
        ws = load_workbook(p).active
        values = [[c.value for c in row] for row in ws.iter_rows()]
        flat = [v for row in values for v in row if v is not None]
        self.assertIn("SHREEJI BOXES", flat)
        self.assertEqual(sum(1 for row in values if row[1] in ("Amit", "Bharat")), 3)

    def test_pdf_written_and_html_has_totals(self):
        p = exports.export_daily_summary_pdf(self.s, self.s.rows, Path(self.tmp.name) / "d.pdf", "SJB")
        self.assertTrue(p.stat().st_size > 1000)
        html = exports.daily_summary_html(self.s, self.s.rows, "SJB", "All entries")
        self.assertIn("not the cash balance", html)
        self.assertIn("7,500", html)
